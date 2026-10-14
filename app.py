from __future__ import annotations

import io
import os
import secrets
import sqlite3
from datetime import datetime, timedelta
from functools import wraps
from pathlib import Path
from urllib.parse import urlsplit
from xml.sax.saxutils import escape

from flask import (
    Flask,
    abort,
    flash,
    g,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook, load_workbook
from werkzeug.security import check_password_hash, generate_password_hash


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_EXCEL = BASE_DIR / "DATABASE GLOBALE TPS ELETTRONICA.xlsx"
APP_VERSION = "1.4.0"


SCHEMA = """
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    username TEXT NOT NULL UNIQUE COLLATE NOCASE,
    full_name TEXT NOT NULL,
    password_hash TEXT NOT NULL,
    role TEXT NOT NULL CHECK (role IN ('admin', 'operatore')),
    active INTEGER NOT NULL DEFAULT 1 CHECK (active IN (0, 1)),
    auth_version INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);

CREATE TABLE IF NOT EXISTS components (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    code TEXT NOT NULL UNIQUE COLLATE NOCASE,
    category TEXT,
    value TEXT,
    description TEXT NOT NULL,
    laboratory TEXT,
    cabinet TEXT,
    door TEXT,
    shelf TEXT,
    drawer TEXT,
    quantity INTEGER NOT NULL DEFAULT 0 CHECK (quantity >= 0),
    reorder_threshold INTEGER NOT NULL DEFAULT 0 CHECK (reorder_threshold >= 0),
    target_quantity INTEGER NOT NULL DEFAULT 0 CHECK (target_quantity >= 0),
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);

CREATE TABLE IF NOT EXISTS movement_batches (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    operation_id TEXT NOT NULL UNIQUE,
    movement_type TEXT NOT NULL CHECK (movement_type IN ('prelievo', 'restituzione')),
    recipient TEXT NOT NULL,
    note TEXT,
    user_id INTEGER NOT NULL REFERENCES users(id),
    operator_snapshot TEXT,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);

CREATE TABLE IF NOT EXISTS movements (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    component_id INTEGER NOT NULL REFERENCES components(id),
    movement_type TEXT NOT NULL CHECK (
        movement_type IN ('prelievo', 'restituzione', 'carico', 'rettifica_carico', 'rettifica_scarico')
    ),
    quantity INTEGER NOT NULL CHECK (quantity > 0),
    delta INTEGER NOT NULL CHECK (delta != 0),
    balance_after INTEGER NOT NULL CHECK (balance_after >= 0),
    recipient TEXT,
    note TEXT,
    user_id INTEGER NOT NULL REFERENCES users(id),
    operator_snapshot TEXT,
    operation_id TEXT,
    batch_id INTEGER REFERENCES movement_batches(id),
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);

CREATE TABLE IF NOT EXISTS reorder_lines (
    component_id INTEGER PRIMARY KEY REFERENCES components(id) ON DELETE CASCADE,
    desired_quantity INTEGER NOT NULL CHECK (desired_quantity > 0),
    set_by INTEGER REFERENCES users(id) ON DELETE SET NULL,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);

CREATE TABLE IF NOT EXISTS app_settings (
    key TEXT PRIMARY KEY,
    value TEXT NOT NULL
);

CREATE INDEX IF NOT EXISTS idx_components_code ON components(code);
CREATE INDEX IF NOT EXISTS idx_components_category ON components(category);
CREATE INDEX IF NOT EXISTS idx_movements_component ON movements(component_id);
CREATE INDEX IF NOT EXISTS idx_movements_created ON movements(created_at DESC);
CREATE INDEX IF NOT EXISTS idx_reorder_lines_updated ON reorder_lines(updated_at DESC);

CREATE TRIGGER IF NOT EXISTS movements_no_update
BEFORE UPDATE ON movements
BEGIN
    SELECT RAISE(ABORT, 'Lo storico movimenti non può essere modificato');
END;

CREATE TRIGGER IF NOT EXISTS movements_no_delete
BEFORE DELETE ON movements
BEGIN
    SELECT RAISE(ABORT, 'Lo storico movimenti non può essere cancellato');
END;

CREATE TRIGGER IF NOT EXISTS movement_batches_no_update
BEFORE UPDATE ON movement_batches
BEGIN
    SELECT RAISE(ABORT, 'I gruppi di movimento non possono essere modificati');
END;

CREATE TRIGGER IF NOT EXISTS movement_batches_no_delete
BEFORE DELETE ON movement_batches
BEGIN
    SELECT RAISE(ABORT, 'I gruppi di movimento non possono essere cancellati');
END;
"""


def create_app(test_config: dict | None = None) -> Flask:
    app = Flask(__name__, instance_relative_config=True)
    Path(app.instance_path).mkdir(parents=True, exist_ok=True)
    app.config.from_mapping(
        SECRET_KEY=load_secret_key(Path(app.instance_path)),
        DATABASE=str(Path(app.instance_path) / "magazzino.db"),
        EXCEL_SOURCE=str(DEFAULT_EXCEL),
        CSRF_ENABLED=True,
        PUBLIC_BASE_URL=os.environ.get("MAGAZZINO_PUBLIC_URL", "").rstrip("/"),
        PERMANENT_SESSION_LIFETIME=timedelta(hours=12),
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=os.environ.get("MAGAZZINO_HTTPS", "0") == "1",
        MAX_CONTENT_LENGTH=1024 * 1024,
    )
    trusted_hosts = os.environ.get("MAGAZZINO_TRUSTED_HOSTS", "").strip()
    if trusted_hosts:
        app.config["TRUSTED_HOSTS"] = [
            host.strip() for host in trusted_hosts.split(",") if host.strip()
        ]
    if test_config:
        app.config.update(test_config)

    app.teardown_appcontext(close_db)
    app.before_request(load_current_user)
    app.before_request(ensure_csrf)
    app.after_request(add_security_headers)
    app.context_processor(inject_helpers)

    register_routes(app)

    with app.app_context():
        initialize_database()

    return app


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        g.db = sqlite3.connect(current_app_config("DATABASE"), timeout=10)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
        g.db.execute("PRAGMA busy_timeout = 10000")
    return g.db


def current_app_config(name: str):
    from flask import current_app

    return current_app.config[name]


def close_db(_error=None) -> None:
    db = g.pop("db", None)
    if db is not None:
        db.close()


def load_secret_key(instance_path: Path) -> str:
    configured = os.environ.get("MAGAZZINO_SECRET_KEY", "")
    if configured:
        return configured
    secret_path = instance_path / "secret.key"
    if secret_path.exists():
        return secret_path.read_text(encoding="utf-8").strip()
    generated = secrets.token_hex(32)
    secret_path.write_text(generated, encoding="utf-8")
    return generated


def initialize_database() -> None:
    db = get_db()
    db.execute("PRAGMA journal_mode = WAL")
    db.execute("PRAGMA synchronous = NORMAL")
    db.executescript(SCHEMA)
    migrate_database_fields(db)
    db.execute(
        """CREATE UNIQUE INDEX IF NOT EXISTS idx_movements_operation_id
           ON movements(operation_id) WHERE operation_id IS NOT NULL"""
    )
    db.execute(
        "CREATE INDEX IF NOT EXISTS idx_movements_batch ON movements(batch_id)"
    )
    count = db.execute("SELECT COUNT(*) FROM components").fetchone()[0]
    if count == 0:
        source = Path(current_app_config("EXCEL_SOURCE"))
        if source.exists():
            import_excel(source, db)
    migrate_reorder_lines(db)
    db.commit()


def migrate_database_fields(db: sqlite3.Connection) -> None:
    user_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(users)").fetchall()
    }
    if "auth_version" not in user_columns:
        db.execute(
            "ALTER TABLE users ADD COLUMN auth_version INTEGER NOT NULL DEFAULT 1"
        )

    columns = {
        row["name"] for row in db.execute("PRAGMA table_info(movements)").fetchall()
    }
    if "operation_id" not in columns:
        db.execute("ALTER TABLE movements ADD COLUMN operation_id TEXT")
    if "batch_id" not in columns:
        db.execute(
            "ALTER TABLE movements ADD COLUMN batch_id INTEGER REFERENCES movement_batches(id)"
        )
    if "operator_snapshot" not in columns:
        db.execute("DROP TRIGGER IF EXISTS movements_no_update")
        db.execute("ALTER TABLE movements ADD COLUMN operator_snapshot TEXT")
    missing_movement_snapshots = db.execute(
        """SELECT 1 FROM movements m
           JOIN users u ON u.id = m.user_id
           WHERE m.operator_snapshot IS NULL LIMIT 1"""
    ).fetchone()
    if missing_movement_snapshots is not None:
        db.execute("DROP TRIGGER IF EXISTS movements_no_update")
        db.execute(
            """UPDATE movements
               SET operator_snapshot = (
                   SELECT full_name FROM users WHERE users.id = movements.user_id
               )
               WHERE operator_snapshot IS NULL"""
        )
    db.executescript(
        """CREATE TRIGGER IF NOT EXISTS movements_no_update
           BEFORE UPDATE ON movements
           BEGIN
               SELECT RAISE(ABORT, 'Lo storico movimenti non può essere modificato');
           END;"""
    )

    batch_columns = {
        row["name"]
        for row in db.execute("PRAGMA table_info(movement_batches)").fetchall()
    }
    if "operator_snapshot" not in batch_columns:
        db.execute("DROP TRIGGER IF EXISTS movement_batches_no_update")
        db.execute(
            "ALTER TABLE movement_batches ADD COLUMN operator_snapshot TEXT"
        )
    missing_batch_snapshots = db.execute(
        """SELECT 1 FROM movement_batches b
           JOIN users u ON u.id = b.user_id
           WHERE b.operator_snapshot IS NULL LIMIT 1"""
    ).fetchone()
    if missing_batch_snapshots is not None:
        db.execute("DROP TRIGGER IF EXISTS movement_batches_no_update")
        db.execute(
            """UPDATE movement_batches
               SET operator_snapshot = (
                   SELECT full_name FROM users
                   WHERE users.id = movement_batches.user_id
               )
               WHERE operator_snapshot IS NULL"""
        )
    db.executescript(
        """CREATE TRIGGER IF NOT EXISTS movement_batches_no_update
           BEFORE UPDATE ON movement_batches
           BEGIN
               SELECT RAISE(ABORT, 'I gruppi di movimento non possono essere modificati');
           END;"""
    )


# Nome mantenuto per compatibilità con integrazioni che richiamavano la migrazione.
def migrate_movement_operation_ids(db: sqlite3.Connection) -> None:
    migrate_database_fields(db)


def migrate_reorder_lines(db: sqlite3.Connection) -> None:
    migration_key = "reorder_lines_migrated_from_targets_v1"
    migrated = db.execute(
        "SELECT 1 FROM app_settings WHERE key = ?", (migration_key,)
    ).fetchone()
    if migrated is not None:
        return
    db.execute(
        """
        INSERT OR IGNORE INTO reorder_lines (component_id, desired_quantity)
        SELECT c.id, c.target_quantity - c.quantity
        FROM components c
        WHERE (c.quantity = 0
               OR (c.reorder_threshold > 0 AND c.quantity <= c.reorder_threshold))
          AND c.target_quantity > c.quantity
          AND (
              NOT EXISTS (
                  SELECT 1 FROM movements m WHERE m.component_id = c.id
              )
              OR c.updated_at > (
                  SELECT MAX(m.created_at)
                  FROM movements m
                  WHERE m.component_id = c.id
              )
          )
        """
    )
    db.execute(
        "INSERT INTO app_settings (key, value) VALUES (?, '1')",
        (migration_key,),
    )


def clean_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def non_negative_int(value, field_name: str) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field_name}: inserire un numero intero") from exc
    if parsed < 0:
        raise ValueError(f"{field_name}: il valore non può essere negativo")
    return parsed


def positive_int(value, field_name: str = "Quantità") -> int:
    parsed = non_negative_int(value, field_name)
    if parsed == 0:
        raise ValueError(f"{field_name}: il valore deve essere maggiore di zero")
    return parsed


def new_operation_id() -> str:
    return secrets.token_urlsafe(24)


def validate_operation_id(value: str | None) -> str:
    operation_id = (value or "").strip()
    if not 20 <= len(operation_id) <= 100 or not all(
        character.isalnum() or character in "-_" for character in operation_id
    ):
        raise ValueError("Operazione non valida. Ricaricare la pagina e riprovare")
    return operation_id


def get_setting(key: str) -> str | None:
    row = get_db().execute(
        "SELECT value FROM app_settings WHERE key = ?", (key,)
    ).fetchone()
    return row["value"] if row is not None else None


def set_setting(key: str, value: str) -> None:
    get_db().execute(
        """INSERT INTO app_settings (key, value) VALUES (?, ?)
           ON CONFLICT(key) DO UPDATE SET value = excluded.value""",
        (key, value),
    )


def normalize_public_base_url(value: str | None) -> str:
    normalized = (value or "").strip().rstrip("/")
    parsed = urlsplit(normalized)
    if (
        parsed.scheme not in {"http", "https"}
        or not parsed.netloc
        or parsed.username
        or parsed.password
        or parsed.query
        or parsed.fragment
        or parsed.path not in {"", "/"}
    ):
        raise ValueError(
            "Inserire un indirizzo completo, ad esempio https://magazzino.scuola.it"
        )
    return normalized


def get_public_base_url() -> str:
    configured = get_setting("public_base_url") or current_app_config("PUBLIC_BASE_URL")
    if not configured:
        return ""
    return normalize_public_base_url(configured)


def component_quick_url(component_code: str) -> str | None:
    base_url = get_public_base_url()
    if not base_url:
        return None
    return base_url + url_for("quick_component", code=component_code)


def begin_inventory_write(db: sqlite3.Connection) -> None:
    try:
        db.execute("BEGIN IMMEDIATE")
    except sqlite3.OperationalError as error:
        if "locked" in str(error).lower() or "busy" in str(error).lower():
            raise ValueError(
                "Il magazzino è momentaneamente occupato. Attendere qualche secondo e riprovare."
            ) from error
        raise


def record_inventory_movement(
    db: sqlite3.Connection,
    *,
    component_id: int,
    movement_type: str,
    quantity: int,
    recipient: str,
    note: str,
    user_id: int,
    operation_id: str,
) -> tuple[sqlite3.Row, bool]:
    begin_inventory_write(db)
    conflicting_batch = db.execute(
        "SELECT 1 FROM movement_batches WHERE operation_id = ?", (operation_id,)
    ).fetchone()
    if conflicting_batch is not None:
        raise ValueError(
            "Identificativo operazione già utilizzato con dati diversi. Ricaricare la pagina."
        )
    existing = db.execute(
        """SELECT m.*, c.code, c.description
           FROM movements m
           JOIN components c ON c.id = m.component_id
           WHERE m.operation_id = ?""",
        (operation_id,),
    ).fetchone()
    if existing is not None:
        expected = (
            component_id,
            movement_type,
            quantity,
            recipient,
            note,
            user_id,
        )
        recorded = (
            existing["component_id"],
            existing["movement_type"],
            existing["quantity"],
            existing["recipient"] or "",
            existing["note"] or "",
            existing["user_id"],
        )
        if recorded != expected:
            raise ValueError(
                "Identificativo operazione già utilizzato con dati diversi. Ricaricare la pagina."
            )
        return existing, True

    component = db.execute(
        "SELECT * FROM components WHERE id = ?", (component_id,)
    ).fetchone()
    if component is None:
        raise ValueError("Componente non trovato")
    operator = db.execute(
        "SELECT full_name FROM users WHERE id = ? AND active = 1", (user_id,)
    ).fetchone()
    if operator is None:
        raise ValueError("L’account non è più attivo. Effettuare nuovamente l’accesso")
    outgoing = movement_type in {"prelievo", "rettifica_scarico"}
    delta = -quantity if outgoing else quantity
    balance = component["quantity"] + delta
    if balance < 0:
        raise ValueError(
            f"Giacenza insufficiente: disponibili {component['quantity']} pezzi"
        )
    db.execute(
        "UPDATE components SET quantity = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
        (balance, component_id),
    )
    cursor = db.execute(
        """
        INSERT INTO movements (
            component_id, movement_type, quantity, delta, balance_after,
            recipient, note, user_id, operator_snapshot, operation_id
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            component_id,
            movement_type,
            quantity,
            delta,
            balance,
            recipient,
            note,
            user_id,
            operator["full_name"],
            operation_id,
        ),
    )
    return db.execute(
        """SELECT m.*, c.code, c.description
           FROM movements m
           JOIN components c ON c.id = m.component_id
           WHERE m.id = ?""",
        (cursor.lastrowid,),
    ).fetchone(), False


def record_inventory_movement_batch(
    db: sqlite3.Connection,
    *,
    movement_type: str,
    lines: list[tuple[int, int]],
    recipient: str,
    note: str,
    user_id: int,
    operation_id: str,
) -> tuple[list[sqlite3.Row], bool]:
    """Registra più righe come un'unica operazione atomica e idempotente."""
    if movement_type not in {"prelievo", "restituzione"}:
        raise ValueError("Tipo di movimento multiplo non valido")
    if not lines:
        raise ValueError("Aggiungere almeno un componente")
    if any(component_id <= 0 or quantity <= 0 for component_id, quantity in lines):
        raise ValueError("Ogni riga deve avere un componente e una quantità valida")

    begin_inventory_write(db)
    conflicting_single = db.execute(
        "SELECT 1 FROM movements WHERE operation_id = ?", (operation_id,)
    ).fetchone()
    if conflicting_single is not None:
        raise ValueError(
            "Identificativo operazione già utilizzato con dati diversi. Ricaricare la pagina."
        )
    existing_batch = db.execute(
        "SELECT * FROM movement_batches WHERE operation_id = ?", (operation_id,)
    ).fetchone()
    if existing_batch is not None:
        existing = db.execute(
            """SELECT m.*, c.code, c.description
               FROM movements m
               JOIN components c ON c.id = m.component_id
               WHERE m.batch_id = ? ORDER BY m.id""",
            (existing_batch["id"],),
        ).fetchall()
        recorded_lines = [
            (movement["component_id"], movement["quantity"])
            for movement in existing
        ]
        if (
            existing_batch["movement_type"] != movement_type
            or existing_batch["recipient"] != recipient
            or (existing_batch["note"] or "") != note
            or existing_batch["user_id"] != user_id
            or recorded_lines != lines
        ):
            raise ValueError(
                "Identificativo operazione già utilizzato con dati diversi. Ricaricare la pagina."
            )
        return list(existing), True

    component_ids = sorted({component_id for component_id, _quantity in lines})
    placeholders = ", ".join("?" for _item in component_ids)
    components = db.execute(
        f"SELECT * FROM components WHERE id IN ({placeholders})", component_ids
    ).fetchall()
    by_id = {component["id"]: component for component in components}
    if len(by_id) != len(component_ids):
        raise ValueError("Uno o più componenti non sono disponibili")
    operator = db.execute(
        "SELECT full_name FROM users WHERE id = ? AND active = 1", (user_id,)
    ).fetchone()
    if operator is None:
        raise ValueError("L’account non è più attivo. Effettuare nuovamente l’accesso")

    totals: dict[int, int] = {}
    for component_id, quantity in lines:
        totals[component_id] = totals.get(component_id, 0) + quantity
    if movement_type == "prelievo":
        for component_id, total in totals.items():
            available = by_id[component_id]["quantity"]
            if total > available:
                raise ValueError(
                    f"Giacenza insufficiente per {by_id[component_id]['code']}: "
                    f"richiesti {total}, disponibili {available}"
                )

    batch_cursor = db.execute(
        """INSERT INTO movement_batches
           (operation_id, movement_type, recipient, note, user_id, operator_snapshot)
           VALUES (?, ?, ?, ?, ?, ?)""",
        (
            operation_id,
            movement_type,
            recipient,
            note,
            user_id,
            operator["full_name"],
        ),
    )
    batch_id = batch_cursor.lastrowid
    balances = {
        component_id: by_id[component_id]["quantity"] for component_id in component_ids
    }
    movement_ids = []
    direction = -1 if movement_type == "prelievo" else 1
    for component_id, quantity in lines:
        delta = direction * quantity
        balances[component_id] += delta
        db.execute(
            """UPDATE components
               SET quantity = ?, updated_at = CURRENT_TIMESTAMP
               WHERE id = ?""",
            (balances[component_id], component_id),
        )
        cursor = db.execute(
            """INSERT INTO movements (
                   component_id, movement_type, quantity, delta, balance_after,
                   recipient, note, user_id, operator_snapshot, batch_id
               ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                component_id,
                movement_type,
                quantity,
                delta,
                balances[component_id],
                recipient,
                note,
                user_id,
                operator["full_name"],
                batch_id,
            ),
        )
        movement_ids.append(cursor.lastrowid)

    result_placeholders = ", ".join("?" for _item in movement_ids)
    recorded = db.execute(
        f"""SELECT m.*, c.code, c.description
            FROM movements m
            JOIN components c ON c.id = m.component_id
            WHERE m.id IN ({result_placeholders}) ORDER BY m.id""",
        movement_ids,
    ).fetchall()
    return list(recorded), False


def import_excel(source: Path, db: sqlite3.Connection) -> int:
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook["COMPONENTI"]
    rows = []
    for values in sheet.iter_rows(min_row=3, values_only=True):
        code = clean_cell(values[0])
        if not code:
            continue
        quantity = non_negative_int(values[9] or 0, f"Quantità di {code}")
        threshold = non_negative_int(values[10] or 0, f"Soglia di {code}")
        rows.append(
            (
                code,
                clean_cell(values[1]),
                clean_cell(values[2]),
                clean_cell(values[3]) or code,
                clean_cell(values[4]),
                clean_cell(values[5]),
                clean_cell(values[6]),
                clean_cell(values[7]),
                clean_cell(values[8]),
                quantity,
                threshold,
                quantity,
            )
        )
    db.executemany(
        """
        INSERT INTO components (
            code, category, value, description, laboratory, cabinet, door,
            shelf, drawer, quantity, reorder_threshold, target_quantity
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        rows,
    )
    return len(rows)


def load_current_user() -> None:
    user_id = session.get("user_id")
    g.user = None
    if user_id is not None:
        g.user = get_db().execute(
            "SELECT * FROM users WHERE id = ? AND active = 1", (user_id,)
        ).fetchone()
        if g.user is None or session.get("auth_version") != g.user["auth_version"]:
            g.user = None
            session.clear()


def ensure_csrf():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(32)
    if request.method == "POST" and current_app_config("CSRF_ENABLED"):
        supplied = request.form.get("csrf_token", "")
        if not secrets.compare_digest(supplied, session["csrf_token"]):
            abort(400, "Token di sicurezza non valido. Ricaricare la pagina.")


def add_security_headers(response):
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    response.headers.setdefault("Referrer-Policy", "same-origin")
    response.headers.setdefault("X-Robots-Tag", "noindex, nofollow")
    if request.endpoint != "static" and (
        request.endpoint in {"login", "setup", "quick_component"}
        or getattr(g, "user", None) is not None
    ):
        response.headers["Cache-Control"] = "no-store"
    return response


def safe_local_next(value: str | None) -> str | None:
    destination = (value or "").strip()
    if not destination or any(character in destination for character in "\\\r\n"):
        return None
    parsed = urlsplit(destination)
    if parsed.scheme or parsed.netloc or not parsed.path.startswith("/"):
        return None
    if parsed.path.startswith("//"):
        return None
    return destination


def inject_helpers() -> dict:
    return {
        "app_version": APP_VERSION,
        "csrf_token": session.get("csrf_token", ""),
        "location_text": location_text,
        "stock_status": stock_status,
    }


def stock_status(component) -> tuple[str, str]:
    quantity = component["quantity"]
    threshold = component["reorder_threshold"]
    target = component["target_quantity"]
    if quantity == 0:
        return "Esaurito", "danger"
    if threshold > 0 and quantity <= threshold:
        return "Da riordinare", "danger"
    if threshold > 0:
        warning_band = max(1, round(max(target - threshold, threshold) * 0.2))
        if quantity <= threshold + warning_band:
            return "Vicino alla soglia", "warning"
    return "Disponibile", "success"


def login_required(view):
    @wraps(view)
    def wrapped(**kwargs):
        if g.user is None:
            return redirect(url_for("login", next=request.full_path))
        return view(**kwargs)

    return wrapped


def admin_required(view):
    @wraps(view)
    def wrapped(**kwargs):
        if g.user is None:
            return redirect(url_for("login"))
        if g.user["role"] != "admin":
            abort(403)
        return view(**kwargs)

    return wrapped


def first_user_missing() -> bool:
    return get_db().execute("SELECT NOT EXISTS(SELECT 1 FROM users)").fetchone()[0] == 1


def register_routes(app: Flask) -> None:
    @app.route("/setup", methods=("GET", "POST"))
    def setup():
        if not first_user_missing():
            return redirect(url_for("login"))
        if request.method == "POST":
            full_name = request.form.get("full_name", "").strip()
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")
            if (
                not full_name
                or not username
                or len(full_name) > 150
                or len(username) > 80
                or not 8 <= len(password) <= 256
            ):
                flash("Compilare tutti i campi; la password deve avere almeno 8 caratteri.", "error")
            else:
                db = get_db()
                try:
                    db.execute("BEGIN IMMEDIATE")
                    if db.execute("SELECT 1 FROM users LIMIT 1").fetchone() is not None:
                        db.rollback()
                        return redirect(url_for("login"))
                    db.execute(
                        "INSERT INTO users (username, full_name, password_hash, role) VALUES (?, ?, ?, 'admin')",
                        (username, full_name, generate_password_hash(password)),
                    )
                    db.commit()
                    flash("Amministratore creato. Ora puoi accedere.", "success")
                    return redirect(url_for("login"))
                except sqlite3.OperationalError as error:
                    db.rollback()
                    if "locked" not in str(error).lower():
                        raise
                    flash(
                        "Configurazione in corso da un’altra sessione. Riprovare tra qualche secondo.",
                        "error",
                    )
        return render_template("setup.html")

    @app.route("/login", methods=("GET", "POST"))
    def login():
        if first_user_missing():
            return redirect(url_for("setup"))
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")
            user = get_db().execute(
                "SELECT * FROM users WHERE username = ? AND active = 1",
                (username[:80],),
            ).fetchone()
            if (
                len(username) > 80
                or len(password) > 256
                or user is None
                or not check_password_hash(user["password_hash"], password)
            ):
                flash("Credenziali non valide.", "error")
            else:
                session.clear()
                session.permanent = True
                session["user_id"] = user["id"]
                session["auth_version"] = user["auth_version"]
                session["csrf_token"] = secrets.token_urlsafe(32)
                destination = safe_local_next(request.args.get("next"))
                if destination is None:
                    destination = url_for("dashboard")
                return redirect(destination)
        return render_template("login.html")

    @app.post("/logout")
    @login_required
    def logout():
        session.clear()
        return redirect(url_for("login"))

    @app.get("/")
    @login_required
    def dashboard():
        db = get_db()
        stats = db.execute(
            """
            SELECT COUNT(*) total,
                   SUM(quantity = 0) exhausted,
                   SUM(reorder_threshold > 0 AND quantity <= reorder_threshold) below_threshold,
                   SUM(reorder_threshold > 0
                       AND quantity > reorder_threshold
                       AND quantity <= reorder_threshold + MAX(
                           1,
                           ROUND(MAX(target_quantity - reorder_threshold, reorder_threshold) * 0.2)
                       )) approaching,
                   SUM(target_quantity = 0) missing_target
            FROM components
            """
        ).fetchone()
        recent = db.execute(
            """
            SELECT m.*, c.code, c.description,
                   COALESCE(m.operator_snapshot, u.full_name) AS operator_name
            FROM movements m
            JOIN components c ON c.id = m.component_id
            JOIN users u ON u.id = m.user_id
            ORDER BY m.id DESC LIMIT 10
            """
        ).fetchall()
        attention = db.execute(
            """
            SELECT * FROM components
            WHERE quantity = 0
               OR (reorder_threshold > 0 AND quantity <= reorder_threshold)
               OR (reorder_threshold > 0
                   AND quantity > reorder_threshold
                   AND quantity <= reorder_threshold + MAX(
                       1,
                       ROUND(MAX(target_quantity - reorder_threshold, reorder_threshold) * 0.2)
                   ))
            ORDER BY (quantity = 0) DESC, quantity - reorder_threshold, code
            LIMIT 12
            """
        ).fetchall()
        return render_template("dashboard.html", stats=stats, recent=recent, attention=attention)

    @app.get("/componenti")
    @login_required
    def components():
        db = get_db()
        query = request.args.get("q", "").strip()
        category = request.args.get("categoria", "").strip()
        status = request.args.get("stato", "").strip()
        clauses, params = [], []
        if query:
            clauses.append("(code LIKE ? OR description LIKE ? OR value LIKE ?)")
            needle = f"%{query}%"
            params.extend((needle, needle, needle))
        if category:
            clauses.append("category = ?")
            params.append(category)
        if status == "esaurito":
            clauses.append("quantity = 0")
        elif status == "riordino":
            clauses.append("target_quantity > quantity AND quantity <= reorder_threshold")
        elif status == "vicino":
            clauses.append(
                """reorder_threshold > 0
                AND quantity > reorder_threshold
                AND quantity <= reorder_threshold + MAX(
                    1,
                    ROUND(MAX(target_quantity - reorder_threshold, reorder_threshold) * 0.2)
                )"""
            )
        elif status == "obiettivo":
            clauses.append("target_quantity = 0")
        sql = "SELECT * FROM components"
        if clauses:
            sql += " WHERE " + " AND ".join(clauses)
        sql += " ORDER BY code LIMIT 1000"
        items = db.execute(sql, params).fetchall()
        categories = db.execute(
            "SELECT DISTINCT category FROM components WHERE category <> '' ORDER BY category"
        ).fetchall()
        return render_template("components.html", components=items, categories=categories)

    @app.get("/componenti/<int:component_id>")
    @login_required
    def component_detail(component_id: int):
        db = get_db()
        component = db.execute("SELECT * FROM components WHERE id = ?", (component_id,)).fetchone()
        if component is None:
            abort(404)
        movements = db.execute(
            """
            SELECT m.*, COALESCE(m.operator_snapshot, u.full_name) AS operator_name
            FROM movements m
            JOIN users u ON u.id = m.user_id
            WHERE m.component_id = ? ORDER BY m.id DESC LIMIT 100
            """,
            (component_id,),
        ).fetchall()
        return render_template(
            "component_detail.html",
            component=component,
            movements=movements,
            quick_url=component_quick_url(component["code"]),
        )

    @app.route("/q/<code>", methods=("GET", "POST"))
    def quick_component(code: str):
        db = get_db()
        component = db.execute(
            "SELECT * FROM components WHERE code = ?", (code,)
        ).fetchone()
        if component is None:
            abort(404)
        if request.method == "POST" and g.user is None:
            return redirect(url_for("login", next=url_for("quick_component", code=code)))
        if request.method == "POST":
            try:
                movement_type = request.form.get("movement_type", "prelievo")
                if movement_type not in {"prelievo", "restituzione"}:
                    raise ValueError("Operazione non consentita dal prelievo rapido")
                quantity = positive_int(request.form.get("quantity"))
                operation_id = validate_operation_id(
                    request.form.get("operation_id")
                )
                destination = request.form.get("destination", "").strip()
                details = request.form.get("note", "").strip()
                if len(destination) > 120:
                    raise ValueError("La destinazione non può superare 120 caratteri")
                if len(details) > 500:
                    raise ValueError("La nota non può superare 500 caratteri")
                note_parts = []
                if destination:
                    note_parts.append(f"Destinazione: {destination}")
                if details:
                    note_parts.append(details)
                movement, duplicate = record_inventory_movement(
                    db,
                    component_id=component["id"],
                    movement_type=movement_type,
                    quantity=quantity,
                    recipient=g.user["full_name"],
                    note=" — ".join(note_parts),
                    user_id=g.user["id"],
                    operation_id=operation_id,
                )
                db.commit()
                if duplicate:
                    flash(
                        "Operazione già registrata: la giacenza non è stata modificata due volte.",
                        "success",
                    )
                elif movement_type == "prelievo":
                    flash(
                        f"Prelevati {quantity} pezzi di {component['code']}. "
                        f"Giacenza rimanente: {movement['balance_after']}.",
                        "success",
                    )
                else:
                    flash(
                        f"Restituiti {quantity} pezzi di {component['code']}. "
                        f"Nuova giacenza: {movement['balance_after']}.",
                        "success",
                    )
                return redirect(url_for("quick_component", code=component["code"]))
            except ValueError as error:
                db.rollback()
                flash(str(error), "error")
                component = db.execute(
                    "SELECT * FROM components WHERE id = ?", (component["id"],)
                ).fetchone()
        recent = []
        if g.user is not None:
            recent = db.execute(
                """SELECT m.*, COALESCE(m.operator_snapshot, u.full_name) AS operator_name
                   FROM movements m
                   JOIN users u ON u.id = m.user_id
                   WHERE m.component_id = ?
                   ORDER BY m.id DESC LIMIT 5""",
                (component["id"],),
            ).fetchall()
        return render_template(
            "quick_component.html",
            component=component,
            recent=recent,
            operation_id=new_operation_id() if g.user is not None else None,
        )

    @app.route("/componenti/<int:component_id>/modifica", methods=("GET", "POST"))
    @admin_required
    def component_edit(component_id: int):
        db = get_db()
        component = db.execute("SELECT * FROM components WHERE id = ?", (component_id,)).fetchone()
        if component is None:
            abort(404)
        if request.method == "POST":
            try:
                threshold = non_negative_int(request.form.get("reorder_threshold"), "Soglia")
                target = non_negative_int(request.form.get("target_quantity"), "Quantità obiettivo")
                if target and threshold > target:
                    raise ValueError("La soglia non può superare la quantità obiettivo")
                description = request.form.get("description", "").strip()
                if not description:
                    raise ValueError("La descrizione è obbligatoria")
                db.execute(
                    """
                    UPDATE components SET description = ?, category = ?, value = ?, laboratory = ?,
                        cabinet = ?, door = ?, shelf = ?, drawer = ?, reorder_threshold = ?,
                        target_quantity = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?
                    """,
                    (
                        description,
                        request.form.get("category", "").strip(),
                        request.form.get("value", "").strip(),
                        request.form.get("laboratory", "").strip(),
                        request.form.get("cabinet", "").strip(),
                        request.form.get("door", "").strip(),
                        request.form.get("shelf", "").strip(),
                        request.form.get("drawer", "").strip(),
                        threshold,
                        target,
                        component_id,
                    ),
                )
                db.commit()
                flash("Dati del componente aggiornati.", "success")
                return redirect(url_for("component_detail", component_id=component_id))
            except ValueError as error:
                flash(str(error), "error")
        return render_template("component_form.html", component=component)

    @app.get("/qr/<code>.svg")
    @login_required
    def component_qr_svg(code: str):
        component = get_db().execute(
            "SELECT code FROM components WHERE code = ?", (code,)
        ).fetchone()
        if component is None:
            abort(404)
        target_url = component_quick_url(component["code"])
        if target_url is None:
            abort(409, "Configurare prima l'indirizzo pubblico per i QR")
        import qrcode
        from qrcode.image.svg import SvgPathImage

        image = qrcode.make(
            target_url,
            image_factory=SvgPathImage,
            box_size=10,
            border=4,
        )
        output = io.BytesIO()
        image.save(output)
        output.seek(0)
        return send_file(
            output,
            mimetype="image/svg+xml",
            download_name=f"qr-{component['code']}.svg",
            max_age=0,
        )

    @app.route("/etichette-qr", methods=("GET", "POST"))
    @admin_required
    def qr_labels():
        db = get_db()
        if request.method == "POST":
            try:
                public_base_url = normalize_public_base_url(
                    request.form.get("public_base_url")
                )
                set_setting("public_base_url", public_base_url)
                db.commit()
                flash("Indirizzo dei QR salvato.", "success")
                return redirect(url_for("qr_labels"))
            except ValueError as error:
                db.rollback()
                flash(str(error), "error")

        query = request.args.get("q", "").strip()
        category = request.args.get("categoria", "").strip()
        laboratory = request.args.get("laboratorio", "").strip()
        clauses, params = [], []
        if query:
            clauses.append("(code LIKE ? OR description LIKE ? OR value LIKE ?)")
            needle = f"%{query}%"
            params.extend((needle, needle, needle))
        if category:
            clauses.append("category = ?")
            params.append(category)
        if laboratory:
            clauses.append("laboratory = ?")
            params.append(laboratory)
        sql = "SELECT * FROM components"
        if clauses:
            sql += " WHERE " + " AND ".join(clauses)
        sql += " ORDER BY code LIMIT 1000"
        items = db.execute(sql, params).fetchall()
        categories = db.execute(
            "SELECT DISTINCT category FROM components WHERE category <> '' ORDER BY category"
        ).fetchall()
        laboratories = db.execute(
            "SELECT DISTINCT laboratory FROM components WHERE laboratory <> '' ORDER BY laboratory"
        ).fetchall()
        public_base_url = get_public_base_url()
        parsed_public_url = urlsplit(public_base_url) if public_base_url else None
        hostname = parsed_public_url.hostname if parsed_public_url else None
        return render_template(
            "qr_labels.html",
            components=items,
            categories=categories,
            laboratories=laboratories,
            public_base_url=public_base_url,
            local_address=hostname in {"localhost", "127.0.0.1", "::1"},
            insecure_address=(
                parsed_public_url is not None and parsed_public_url.scheme != "https"
            ),
        )

    @app.post("/etichette-qr.pdf")
    @admin_required
    def qr_labels_pdf():
        db = get_db()
        try:
            base_url = get_public_base_url()
            if not base_url:
                raise ValueError("Configurare prima l'indirizzo pubblico dei QR")
            try:
                component_ids = sorted(
                    {int(value) for value in request.form.getlist("selected_ids")}
                )
            except ValueError as exc:
                raise ValueError("Selezione dei componenti non valida") from exc
            if not component_ids:
                raise ValueError("Selezionare almeno un componente")
            placeholders = ",".join("?" for _ in component_ids)
            items = db.execute(
                f"SELECT * FROM components WHERE id IN ({placeholders}) ORDER BY code",
                component_ids,
            ).fetchall()
            if len(items) != len(component_ids):
                raise ValueError("Uno o più componenti non sono più disponibili")
            output = build_qr_labels_pdf(items, base_url)
            stamp = datetime.now().strftime("%Y-%m-%d")
            return send_file(
                output,
                as_attachment=True,
                download_name=f"etichette-qr-{stamp}.pdf",
                mimetype="application/pdf",
            )
        except ValueError as error:
            flash(str(error), "error")
            return redirect(url_for("qr_labels"))

    @app.get("/componenti/<int:component_id>/etichetta.pdf")
    @admin_required
    def component_qr_label_pdf(component_id: int):
        component = get_db().execute(
            "SELECT * FROM components WHERE id = ?", (component_id,)
        ).fetchone()
        if component is None:
            abort(404)
        base_url = get_public_base_url()
        if not base_url:
            flash("Configurare prima l'indirizzo pubblico dei QR.", "error")
            return redirect(url_for("qr_labels"))
        output = build_qr_labels_pdf([component], base_url)
        return send_file(
            output,
            as_attachment=True,
            download_name=f"etichetta-{component['code']}.pdf",
            mimetype="application/pdf",
        )

    @app.route("/movimenti/nuovo", methods=("GET", "POST"))
    @login_required
    def movement_new():
        db = get_db()
        selected_id = request.args.get("componente", type=int)
        if request.method == "POST":
            selected_id = request.form.get("component_id", type=int)
            movement_type = request.form.get("movement_type", "")
            recipient = request.form.get("recipient", "").strip()
            note = request.form.get("note", "").strip()
            try:
                if len(recipient) > 200:
                    raise ValueError("Il destinatario non può superare 200 caratteri")
                if len(note) > 500:
                    raise ValueError("La nota non può superare 500 caratteri")
                if movement_type not in {
                    "prelievo", "restituzione", "carico", "rettifica_carico", "rettifica_scarico"
                }:
                    raise ValueError("Tipo di movimento non valido")
                if g.user["role"] != "admin" and movement_type not in {
                    "prelievo",
                    "restituzione",
                }:
                    raise ValueError(
                        "I docenti possono registrare solo prelievi e restituzioni"
                    )
                quantity = positive_int(request.form.get("quantity"))
                if movement_type == "prelievo" and not recipient:
                    raise ValueError("Indicare chi preleva il materiale")
                operation_id = validate_operation_id(
                    request.form.get("operation_id") or new_operation_id()
                )
                movement, duplicate = record_inventory_movement(
                    db,
                    component_id=selected_id,
                    movement_type=movement_type,
                    quantity=quantity,
                    recipient=recipient,
                    note=note,
                    user_id=g.user["id"],
                    operation_id=operation_id,
                )
                db.commit()
                if duplicate:
                    flash(
                        "Operazione già registrata: la giacenza non è stata modificata due volte.",
                        "success",
                    )
                else:
                    flash(
                        f"Movimento registrato. Nuova giacenza: {movement['balance_after']}.",
                        "success",
                    )
                return redirect(url_for("component_detail", component_id=selected_id))
            except ValueError as error:
                db.rollback()
                flash(str(error), "error")
        components_list = db.execute(
            "SELECT id, code, description, quantity FROM components ORDER BY code"
        ).fetchall()
        return render_template(
            "movement_form.html",
            components=components_list,
            selected_id=selected_id,
            operation_id=new_operation_id(),
        )

    @app.route("/movimenti/multipli", methods=("GET", "POST"))
    @login_required
    def movement_batch():
        db = get_db()
        submitted_lines = []
        if request.method == "POST":
            component_values = request.form.getlist("component_id")
            quantity_values = request.form.getlist("quantity")
            row_count = max(len(component_values), len(quantity_values))
            for index in range(row_count):
                submitted_lines.append(
                    {
                        "component_id": (
                            component_values[index] if index < len(component_values) else ""
                        ),
                        "quantity": (
                            quantity_values[index] if index < len(quantity_values) else ""
                        ),
                    }
                )

            try:
                if row_count > 100:
                    raise ValueError("È possibile registrare al massimo 100 righe alla volta")
                movement_type = request.form.get("movement_type", "")
                if movement_type not in {"prelievo", "restituzione"}:
                    raise ValueError("Selezionare prelievo o riconsegna")

                parsed_lines = []
                for index, line in enumerate(submitted_lines, start=1):
                    raw_component = str(line["component_id"]).strip()
                    raw_quantity = str(line["quantity"]).strip()
                    if not raw_component and not raw_quantity:
                        continue
                    if not raw_component or not raw_quantity:
                        raise ValueError(
                            f"Riga {index}: scegliere il componente e indicare la quantità"
                        )
                    try:
                        component_id = int(raw_component)
                    except ValueError as error:
                        raise ValueError(f"Riga {index}: componente non valido") from error
                    quantity = positive_int(raw_quantity, f"Quantità alla riga {index}")
                    parsed_lines.append((component_id, quantity))

                destination = request.form.get("destination", "").strip()
                details = request.form.get("note", "").strip()
                if len(destination) > 200:
                    raise ValueError("La destinazione non può superare 200 caratteri")
                if len(details) > 500:
                    raise ValueError("Le note non possono superare 500 caratteri")
                note_parts = []
                if destination:
                    note_parts.append(f"Destinazione: {destination}")
                if details:
                    note_parts.append(details)

                operation_id = validate_operation_id(request.form.get("operation_id"))
                recorded, duplicate = record_inventory_movement_batch(
                    db,
                    movement_type=movement_type,
                    lines=parsed_lines,
                    recipient=g.user["full_name"],
                    note=" — ".join(note_parts),
                    user_id=g.user["id"],
                    operation_id=operation_id,
                )
                db.commit()
                if duplicate:
                    flash(
                        "Operazione già registrata: le giacenze non sono state modificate due volte.",
                        "success",
                    )
                else:
                    pieces = sum(row["quantity"] for row in recorded)
                    action = "Prelievo" if movement_type == "prelievo" else "Riconsegna"
                    flash(
                        f"{action} registrato: {pieces} pezzi in {len(recorded)} righe.",
                        "success",
                    )
                return redirect(url_for("movement_batch"))
            except ValueError as error:
                db.rollback()
                flash(str(error), "error")

        selected_id = request.args.get("componente", type=int)
        if not submitted_lines:
            submitted_lines = [
                {"component_id": selected_id or "", "quantity": ""},
            ]
        components_list = db.execute(
            "SELECT id, code, description, quantity FROM components ORDER BY code"
        ).fetchall()
        return render_template(
            "movement_batch_form.html",
            components=components_list,
            submitted_lines=submitted_lines,
            operation_id=new_operation_id(),
        )

    @app.get("/movimenti")
    @login_required
    def movements():
        rows = get_db().execute(
            """
            SELECT m.*, c.code, c.description,
                   COALESCE(m.operator_snapshot, u.full_name) AS operator_name
            FROM movements m JOIN components c ON c.id = m.component_id
            JOIN users u ON u.id = m.user_id
            ORDER BY m.id DESC LIMIT 1000
            """
        ).fetchall()
        return render_template("movements.html", movements=rows)

    @app.get("/riordino")
    @login_required
    def reorder_report():
        db = get_db()
        return render_template(
            "reorder.html",
            items=get_reorder_items(),
            approaching=get_approaching_items(),
            module_count=db.execute("SELECT COUNT(*) FROM reorder_lines").fetchone()[0],
        )

    @app.post("/riordino/quantita")
    @admin_required
    def reorder_quantities():
        db = get_db()
        action = request.form.get("action", "")
        destination = "reorder_report"
        try:
            if action == "bulk":
                raw_ids = request.form.getlist("selected_ids")
                try:
                    component_ids = sorted({int(value) for value in raw_ids})
                except ValueError as exc:
                    raise ValueError("Selezione dei componenti non valida") from exc
                if not component_ids:
                    raise ValueError("Selezionare almeno un componente")
                order_quantity = positive_int(
                    request.form.get("bulk_order_quantity"), "Quantità da ordinare"
                )
                db.execute("BEGIN IMMEDIATE")
                components_to_update = components_by_ids(db, component_ids)
                if len(components_to_update) != len(component_ids):
                    raise ValueError("Uno o più componenti non sono più presenti nella distinta")
                db.executemany(
                    """
                    INSERT INTO reorder_lines (component_id, desired_quantity, set_by)
                    VALUES (?, ?, ?)
                    ON CONFLICT(component_id) DO UPDATE SET
                        desired_quantity = excluded.desired_quantity,
                        set_by = excluded.set_by,
                        updated_at = CURRENT_TIMESTAMP
                    """,
                    [
                        (component["id"], order_quantity, g.user["id"])
                        for component in components_to_update
                    ],
                )
                db.commit()
                flash(
                    f"Aggiunti al modulo {len(component_ids)} componenti, quantità {order_quantity} ciascuno.",
                    "success",
                )
                destination = "order_form"
            elif action.startswith("single:"):
                try:
                    component_id = int(action.split(":", 1)[1])
                except (IndexError, ValueError) as exc:
                    raise ValueError("Componente non valido") from exc
                order_quantity = positive_int(
                    request.form.get(f"order_quantity_{component_id}"),
                    "Quantità da ordinare",
                )
                db.execute("BEGIN IMMEDIATE")
                components_to_update = components_by_ids(db, [component_id])
                if not components_to_update:
                    raise ValueError("Il componente non è più presente nella distinta")
                component = components_to_update[0]
                db.execute(
                    """
                    INSERT INTO reorder_lines (component_id, desired_quantity, set_by)
                    VALUES (?, ?, ?)
                    ON CONFLICT(component_id) DO UPDATE SET
                        desired_quantity = excluded.desired_quantity,
                        set_by = excluded.set_by,
                        updated_at = CURRENT_TIMESTAMP
                    """,
                    (component_id, order_quantity, g.user["id"]),
                )
                db.commit()
                flash(
                    f"{component['code']} aggiunto al modulo con quantità {order_quantity}.",
                    "success",
                )
                destination = "order_form"
            else:
                raise ValueError("Operazione non valida")
        except ValueError as error:
            db.rollback()
            flash(str(error), "error")
        return redirect(url_for(destination))

    @app.get("/modulo-riordino")
    @login_required
    def order_form():
        items = get_order_form_items()
        return render_template(
            "order_form.html",
            items=items,
            total_pieces=sum(item["desired_quantity"] for item in items),
            generated_at=datetime.now(),
        )

    @app.post("/modulo-riordino/<int:component_id>/quantita")
    @admin_required
    def order_form_quantity(component_id: int):
        db = get_db()
        try:
            desired_quantity = positive_int(
                request.form.get("desired_quantity"), "Quantità desiderata"
            )
            cursor = db.execute(
                """
                UPDATE reorder_lines
                SET desired_quantity = ?, set_by = ?, updated_at = CURRENT_TIMESTAMP
                WHERE component_id = ?
                """,
                (desired_quantity, g.user["id"], component_id),
            )
            if cursor.rowcount == 0:
                raise ValueError("Riga del modulo non trovata")
            db.commit()
            flash("Quantità del modulo aggiornata.", "success")
        except ValueError as error:
            db.rollback()
            flash(str(error), "error")
        return redirect(url_for("order_form"))

    @app.post("/modulo-riordino/<int:component_id>/rimuovi")
    @admin_required
    def order_form_remove(component_id: int):
        db = get_db()
        cursor = db.execute(
            "DELETE FROM reorder_lines WHERE component_id = ?", (component_id,)
        )
        db.commit()
        if cursor.rowcount:
            flash("Componente rimosso dal modulo e riportato nella distinta.", "success")
        else:
            flash("Riga del modulo non trovata.", "error")
        return redirect(url_for("order_form"))

    @app.get("/riordino.xlsx")
    @login_required
    def reorder_excel():
        items = get_reorder_items()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Distinta da completare"
        sheet.append(
            ["Codice", "Descrizione", "Categoria", "Giacenza", "Soglia", "Obiettivo", "Quantità suggerita", "Ubicazione"]
        )
        for item in items:
            target = item["target_quantity"] if item["target_quantity"] > 0 else "DA DEFINIRE"
            to_order = item["to_order"] if item["to_order"] is not None else "DA DEFINIRE"
            sheet.append(
                [
                    item["code"], item["description"], item["category"], item["quantity"],
                    item["reorder_threshold"], target, to_order,
                    location_text(item),
                ]
            )
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                55, max(len(str(cell.value or "")) for cell in column) + 2
            )
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        stamp = datetime.now().strftime("%Y-%m-%d")
        return send_file(
            output,
            as_attachment=True,
            download_name=f"distinta-da-completare-{stamp}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.get("/riordino.pdf")
    @login_required
    def reorder_pdf():
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        items = get_reorder_items()
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
            topMargin=10 * mm, bottomMargin=10 * mm,
        )
        styles = getSampleStyleSheet()
        story = [
            Paragraph("ITIS Cannizzaro Colleferro — Distinta da completare", styles["Title"]),
            Paragraph(datetime.now().strftime("Generata il %d/%m/%Y alle %H:%M"), styles["Normal"]),
            Spacer(1, 5 * mm),
        ]
        data = [["Codice", "Descrizione", "Giac.", "Soglia", "Obiettivo", "Suggerita"]]
        for item in items:
            target = item["target_quantity"] if item["target_quantity"] > 0 else "Da definire"
            to_order = item["to_order"] if item["to_order"] is not None else "Da definire"
            data.append(
                [item["code"], Paragraph(escape(str(item["description"])), styles["BodyText"]), item["quantity"],
                 item["reorder_threshold"], target, to_order]
            )
        if not items:
            data.append(["—", "Nessun componente da ordinare", "", "", "", ""])
        table = Table(data, colWidths=[28 * mm, 130 * mm, 20 * mm, 22 * mm, 24 * mm, 22 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173d6b")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef3f8")]),
                    ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.append(table)
        doc.build(story)
        output.seek(0)
        stamp = datetime.now().strftime("%Y-%m-%d")
        return send_file(output, as_attachment=True, download_name=f"distinta-da-completare-{stamp}.pdf", mimetype="application/pdf")

    @app.get("/modulo-riordino.xlsx")
    @login_required
    def order_form_excel():
        items = get_order_form_items()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Modulo riordino"
        sheet.append(
            ["Codice", "Descrizione", "Categoria", "Giacenza", "Quantità desiderata", "Ubicazione"]
        )
        for item in items:
            sheet.append(
                [
                    item["code"], item["description"], item["category"], item["quantity"],
                    item["desired_quantity"], location_text(item),
                ]
            )
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                55, max(len(str(cell.value or "")) for cell in column) + 2
            )
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        stamp = datetime.now().strftime("%Y-%m-%d")
        return send_file(
            output,
            as_attachment=True,
            download_name=f"modulo-riordino-{stamp}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.get("/modulo-riordino.pdf")
    @login_required
    def order_form_pdf():
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        items = get_order_form_items()
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
            topMargin=10 * mm, bottomMargin=10 * mm,
        )
        styles = getSampleStyleSheet()
        total_pieces = sum(item["desired_quantity"] for item in items)
        story = [
            Paragraph("ITIS Cannizzaro Colleferro — Modulo riordino", styles["Title"]),
            Paragraph(
                datetime.now().strftime("Generato il %d/%m/%Y alle %H:%M")
                + f" — Compilato da: {escape(str(g.user['full_name']))} — Righe: {len(items)}",
                styles["Normal"],
            ),
            Spacer(1, 5 * mm),
        ]
        data = [["Codice", "Descrizione", "Categoria", "Giac.", "Quantità desiderata", "Ubicazione"]]
        for item in items:
            data.append(
                [
                    item["code"], Paragraph(escape(str(item["description"])), styles["BodyText"]),
                    item["category"], item["quantity"], item["desired_quantity"],
                    location_text(item),
                ]
            )
        if not items:
            data.append(["—", "Modulo vuoto", "", "", "", ""])
        else:
            data.append(["", "", "", "Totale pezzi", total_pieces, ""])
        table = Table(
            data,
            colWidths=[27 * mm, 100 * mm, 35 * mm, 18 * mm, 35 * mm, 55 * mm],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173d6b")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (3, -1), (4, -1), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef3f8")]),
                    ("ALIGN", (3, 1), (4, -1), "RIGHT"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.append(table)
        story.extend(
            [
                Spacer(1, 10 * mm),
                Paragraph("Note: _________________________________________________________________", styles["Normal"]),
                Spacer(1, 10 * mm),
                Paragraph("Firma del responsabile: ______________________________________________", styles["Normal"]),
            ]
        )
        doc.build(story)
        output.seek(0)
        stamp = datetime.now().strftime("%Y-%m-%d")
        return send_file(
            output,
            as_attachment=True,
            download_name=f"modulo-riordino-{stamp}.pdf",
            mimetype="application/pdf",
        )

    @app.route("/utenti", methods=("GET", "POST"))
    @admin_required
    def users():
        db = get_db()
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            full_name = request.form.get("full_name", "").strip()
            password = request.form.get("password", "")
            role = request.form.get("role", "operatore")
            try:
                if (
                    not username
                    or not full_name
                    or len(username) > 80
                    or len(full_name) > 150
                    or not 8 <= len(password) <= 256
                ):
                    raise ValueError("Compilare tutti i campi; password minima di 8 caratteri")
                if role not in {"admin", "operatore"}:
                    raise ValueError("Ruolo non valido")
                db.execute(
                    "INSERT INTO users (username, full_name, password_hash, role) VALUES (?, ?, ?, ?)",
                    (username, full_name, generate_password_hash(password), role),
                )
                db.commit()
                flash("Utente creato.", "success")
                return redirect(url_for("users"))
            except sqlite3.IntegrityError:
                flash("Nome utente già utilizzato.", "error")
            except ValueError as error:
                flash(str(error), "error")
        rows = db.execute("SELECT * FROM users ORDER BY active DESC, full_name").fetchall()
        return render_template("users.html", users=rows)

    @app.post("/utenti/<int:user_id>/stato")
    @admin_required
    def user_toggle(user_id: int):
        db = get_db()
        user = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if user is None:
            abort(404)
        if user_id == g.user["id"]:
            flash("Non puoi disattivare il tuo account.", "error")
        else:
            db.execute(
                """UPDATE users
                   SET active = CASE active WHEN 1 THEN 0 ELSE 1 END,
                       auth_version = auth_version + 1
                   WHERE id = ?""",
                (user_id,),
            )
            db.commit()
            flash("Stato utente aggiornato.", "success")
        return redirect(url_for("users"))

    @app.route("/utenti/<int:user_id>/modifica", methods=("GET", "POST"))
    @admin_required
    def user_edit(user_id: int):
        db = get_db()
        user = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if user is None:
            abort(404)

        if request.method == "POST":
            username = request.form.get("username", "").strip()
            full_name = request.form.get("full_name", "").strip()
            password = request.form.get("password", "")
            role = request.form.get("role", "")
            active = 1 if request.form.get("active") == "1" else 0
            try:
                if not username or not full_name:
                    raise ValueError("Nome, cognome e nome utente sono obbligatori.")
                if len(username) > 80 or len(full_name) > 150:
                    raise ValueError("Nome o nome utente troppo lungo.")
                if role not in {"admin", "operatore"}:
                    raise ValueError("Ruolo non valido.")
                if password and not 8 <= len(password) <= 256:
                    raise ValueError(
                        "La nuova password deve avere tra 8 e 256 caratteri."
                    )
                if user_id == g.user["id"] and (role != "admin" or active != 1):
                    raise ValueError(
                        "Non puoi rimuovere i privilegi o disattivare l’account che stai usando."
                    )

                fields = [
                    "username = ?",
                    "full_name = ?",
                    "role = ?",
                    "active = ?",
                ]
                values = [username, full_name, role, active]
                revoke_sessions = bool(password) or role != user["role"] or active != user["active"]
                if revoke_sessions:
                    fields.append("auth_version = auth_version + 1")
                if password:
                    fields.append("password_hash = ?")
                    values.append(generate_password_hash(password))
                values.append(user_id)
                db.execute(
                    f"UPDATE users SET {', '.join(fields)} WHERE id = ?",
                    values,
                )
                db.commit()
                if user_id == g.user["id"] and revoke_sessions:
                    session["auth_version"] = db.execute(
                        "SELECT auth_version FROM users WHERE id = ?", (user_id,)
                    ).fetchone()[0]
                flash("Utente aggiornato.", "success")
                return redirect(url_for("users"))
            except sqlite3.IntegrityError:
                db.rollback()
                flash("Nome utente già utilizzato.", "error")
            except ValueError as error:
                flash(str(error), "error")

        return render_template("user_form.html", user=user)

    @app.post("/utenti/<int:user_id>/elimina")
    @admin_required
    def user_delete(user_id: int):
        db = get_db()
        user = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if user is None:
            abort(404)
        if user_id == g.user["id"]:
            flash("Non puoi eliminare l’account che stai usando.", "error")
            return redirect(url_for("users"))

        # Serializza il controllo e la cancellazione: un movimento registrato
        # nello stesso istante non deve mai far perdere il riferimento all'utente.
        db.execute("BEGIN IMMEDIATE")
        movement_count = db.execute(
            "SELECT COUNT(*) FROM movements WHERE user_id = ?", (user_id,)
        ).fetchone()[0]
        if movement_count:
            db.execute(
                """UPDATE users
                   SET active = 0, auth_version = auth_version + 1
                   WHERE id = ?""",
                (user_id,),
            )
            db.commit()
            flash(
                "L’account ha movimenti registrati: è stato disattivato e lo storico è stato conservato.",
                "success",
            )
        else:
            db.execute("DELETE FROM users WHERE id = ?", (user_id,))
            db.commit()
            flash("Utente eliminato.", "success")
        return redirect(url_for("users"))


def get_reorder_items():
    return get_db().execute(
        """
        SELECT c.*,
               CASE
                   WHEN c.target_quantity > c.quantity THEN c.target_quantity - c.quantity
                   ELSE NULL
               END AS to_order
        FROM components c
        WHERE (c.quantity = 0
               OR (c.reorder_threshold > 0 AND c.quantity <= c.reorder_threshold))
          AND NOT EXISTS (
              SELECT 1 FROM reorder_lines r WHERE r.component_id = c.id
          )
        ORDER BY (c.target_quantity = 0) DESC, c.category, c.code
        """
    ).fetchall()


def get_order_form_items():
    return get_db().execute(
        """
        SELECT c.*, r.desired_quantity, r.created_at AS order_created_at,
               r.updated_at AS order_updated_at, u.full_name AS set_by_name
        FROM reorder_lines r
        JOIN components c ON c.id = r.component_id
        LEFT JOIN users u ON u.id = r.set_by
        ORDER BY c.category, c.code
        """
    ).fetchall()


def components_by_ids(db: sqlite3.Connection, component_ids: list[int]):
    if not component_ids:
        return []
    placeholders = ",".join("?" for _ in component_ids)
    return db.execute(
        f"""
        SELECT c.id, c.code, c.quantity, c.reorder_threshold
        FROM components c
        WHERE c.id IN ({placeholders})
          AND (c.quantity = 0 OR (c.reorder_threshold > 0 AND c.quantity <= c.reorder_threshold))
          AND NOT EXISTS (
              SELECT 1 FROM reorder_lines r WHERE r.component_id = c.id
          )
        ORDER BY c.id
        """,
        component_ids,
    ).fetchall()


def get_approaching_items():
    return get_db().execute(
        """
        SELECT *, quantity - reorder_threshold AS units_to_threshold
        FROM components
        WHERE reorder_threshold > 0
          AND quantity > reorder_threshold
          AND quantity <= reorder_threshold + MAX(
              1,
              ROUND(MAX(target_quantity - reorder_threshold, reorder_threshold) * 0.2)
          )
        ORDER BY units_to_threshold, category, code
        """
    ).fetchall()


def location_text(component) -> str:
    parts = [
        component["laboratory"],
        f"Arm. {component['cabinet']}" if component["cabinet"] else "",
        f"Anta {component['door']}" if component["door"] else "",
        f"Scaff. {component['shelf']}" if component["shelf"] else "",
        f"Cass. {component['drawer']}" if component["drawer"] else "",
    ]
    return " · ".join(part for part in parts if part)


def build_qr_labels_pdf(items, base_url: str) -> io.BytesIO:
    from reportlab.graphics.barcode.qr import QrCodeWidget
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    def shortened(value, limit: int) -> str:
        text = str(value or "").strip()
        return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"

    def qr_drawing(value: str) -> Drawing:
        widget = QrCodeWidget(value)
        x1, y1, x2, y2 = widget.getBounds()
        width, height = x2 - x1, y2 - y1
        size = 27 * mm
        scale = min(size / width, size / height)
        drawing = Drawing(
            size,
            size,
            transform=[scale, 0, 0, scale, -x1 * scale, -y1 * scale],
        )
        drawing.add(widget)
        return drawing

    code_style = ParagraphStyle(
        "LabelCode",
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=10,
        alignment=TA_LEFT,
        spaceAfter=2,
    )
    detail_style = ParagraphStyle(
        "LabelDetail",
        fontName="Helvetica",
        fontSize=6.5,
        leading=7.5,
        alignment=TA_LEFT,
    )
    labels = []
    for item in items:
        quick_url = base_url.rstrip("/") + url_for(
            "quick_component", code=item["code"]
        )
        details = [
            Paragraph(escape(str(item["code"])), code_style),
            Paragraph(escape(shortened(item["description"], 72)), detail_style),
        ]
        if item["value"]:
            details.append(
                Paragraph(
                    f"Valore: {escape(shortened(item['value'], 32))}", detail_style
                )
            )
        location = location_text(item)
        if location:
            details.append(
                Paragraph(
                    f"Posizione: {escape(shortened(location, 62))}", detail_style
                )
            )
        label = Table(
            [[qr_drawing(quick_url), details]],
            colWidths=[29 * mm, 36 * mm],
            rowHeights=[34 * mm],
        )
        label.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 1.5 * mm),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 1.5 * mm),
                    ("TOPPADDING", (0, 0), (-1, -1), 1 * mm),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 1 * mm),
                ]
            )
        )
        labels.append(label)

    rows = [labels[index : index + 3] for index in range(0, len(labels), 3)]
    if rows and len(rows[-1]) < 3:
        rows[-1].extend([""] * (3 - len(rows[-1])))
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=6 * mm,
        rightMargin=6 * mm,
        topMargin=7 * mm,
        bottomMargin=7 * mm,
    )
    grid = Table(
        rows,
        colWidths=[66 * mm] * 3,
        rowHeights=[35 * mm] * len(rows),
    )
    grid.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#aab4bf")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )
    document.build([grid])
    output.seek(0)
    return output


app = create_app()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "3000")), debug=False)

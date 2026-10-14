import sqlite3
import tempfile
import unittest
import xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

from app import create_app, get_db, migrate_reorder_lines


class WarehouseTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db_path = str(Path(self.temp.name) / "test.db")
        self.app = create_app(
            {
                "TESTING": True,
                "DATABASE": self.db_path,
                "EXCEL_SOURCE": str(Path(self.temp.name) / "missing.xlsx"),
                "CSRF_ENABLED": False,
                "SECRET_KEY": "test",
            }
        )
        with self.app.app_context():
            db = get_db()
            db.execute(
                "INSERT INTO users (username, full_name, password_hash, role) VALUES (?, ?, ?, ?)",
                ("admin", "Admin Test", generate_password_hash("password-test"), "admin"),
            )
            db.execute(
                """
                INSERT INTO components
                    (code, description, quantity, reorder_threshold, target_quantity)
                VALUES ('TPS0001', 'Resistore test', 10, 3, 10)
                """
            )
            db.commit()
        self.client = self.app.test_client()
        self.client.post("/login", data={"username": "admin", "password": "password-test"})

    def create_operator(
        self,
        username: str = "docente",
        full_name: str = "Prof.ssa Maria Rossi",
    ) -> int:
        with self.app.app_context():
            db = get_db()
            cursor = db.execute(
                """INSERT INTO users (username, full_name, password_hash, role)
                   VALUES (?, ?, ?, 'operatore')""",
                (username, full_name, generate_password_hash("password-test")),
            )
            db.commit()
            return cursor.lastrowid

    def use_account(self, user_id: int) -> None:
        with self.client.session_transaction() as session_data:
            session_data.clear()
            session_data["user_id"] = user_id
            session_data["auth_version"] = 1

    def tearDown(self):
        self.temp.cleanup()

    def test_withdrawal_updates_stock_and_audit(self):
        response = self.client.post(
            "/movimenti/nuovo",
            data={
                "component_id": "1",
                "movement_type": "prelievo",
                "quantity": "4",
                "recipient": "Prof. Rossi",
                "note": "Classe 4A",
            },
        )
        self.assertEqual(response.status_code, 302)
        with self.app.app_context():
            db = get_db()
            self.assertEqual(db.execute("SELECT quantity FROM components").fetchone()[0], 6)
            movement = db.execute("SELECT * FROM movements").fetchone()
            self.assertEqual(movement["delta"], -4)
            self.assertEqual(movement["recipient"], "Prof. Rossi")

    def test_stock_cannot_become_negative(self):
        response = self.client.post(
            "/movimenti/nuovo",
            data={
                "component_id": "1",
                "movement_type": "prelievo",
                "quantity": "11",
                "recipient": "Prof. Rossi",
            },
            follow_redirects=True,
        )
        self.assertIn("Giacenza insufficiente".encode(), response.data)
        with self.app.app_context():
            self.assertEqual(get_db().execute("SELECT quantity FROM components").fetchone()[0], 10)

    def test_quick_page_is_public_but_hides_form_and_movement_history(self):
        self.client.post(
            "/movimenti/nuovo",
            data={
                "component_id": "1",
                "movement_type": "prelievo",
                "quantity": "1",
                "recipient": "Persona riservata",
                "note": "Nota riservata",
            },
        )
        with self.client.session_transaction() as session_data:
            session_data.clear()

        public_page = self.client.get("/q/TPS0001")

        self.assertEqual(public_page.status_code, 200)
        self.assertIn(b"TPS0001", public_page.data)
        self.assertIn(b"Resistore test", public_page.data)
        self.assertIn(b">9<", public_page.data)
        self.assertIn(b"/login?next=/q/TPS0001", public_page.data)
        self.assertIn(b"Accedi", public_page.data)
        self.assertNotIn(b"<form", public_page.data)
        self.assertNotIn(b'name="operation_id"', public_page.data)
        self.assertNotIn(b"Ultime operazioni", public_page.data)
        self.assertNotIn(b"Persona riservata", public_page.data)
        self.assertNotIn(b"Nota riservata", public_page.data)
        self.assertNotIn(b">Dashboard<", public_page.data)

    def test_anonymous_quick_post_does_not_change_stock_and_requires_login(self):
        with self.client.session_transaction() as session_data:
            session_data.clear()

        response = self.client.post(
            "/q/TPS0001",
            data={
                "movement_type": "prelievo",
                "quantity": "3",
                "operation_id": "anonymous-attempt-test-0001",
            },
        )

        self.assertEqual(response.status_code, 302)
        redirect_url = urlsplit(response.headers["Location"])
        self.assertEqual(redirect_url.path, "/login")
        self.assertEqual(
            parse_qs(redirect_url.query)["next"][0].rstrip("?"),
            "/q/TPS0001",
        )
        with self.app.app_context():
            db = get_db()
            self.assertEqual(
                db.execute("SELECT quantity FROM components WHERE id = 1").fetchone()[0],
                10,
            )
            self.assertEqual(db.execute("SELECT COUNT(*) FROM movements").fetchone()[0], 0)

    def test_login_next_returns_to_scanned_component_and_enables_withdrawal(self):
        with self.client.session_transaction() as session_data:
            session_data.clear()

        login = self.client.post(
            "/login?next=/q/TPS0001",
            data={"username": "admin", "password": "password-test"},
        )
        self.assertEqual(login.status_code, 302)
        self.assertTrue(login.headers["Location"].endswith("/q/TPS0001"))

        authenticated_page = self.client.get(login.headers["Location"])
        self.assertEqual(authenticated_page.status_code, 200)
        self.assertIn("Prelievo rapido".encode(), authenticated_page.data)
        self.assertIn(b"TPS0001", authenticated_page.data)
        self.assertIn(b"Resistore test", authenticated_page.data)
        self.assertIn(b'name="operation_id"', authenticated_page.data)
        self.assertIn(b">Dashboard<", authenticated_page.data)

    def test_quick_withdrawal_uses_operator_identity_note_and_updates_stock(self):
        operator_id = self.create_operator()
        self.use_account(operator_id)

        response = self.client.post(
            "/q/TPS0001",
            data={
                "movement_type": "prelievo",
                "quantity": "3",
                "destination": "Classe 4A Elettronica",
                "note": "Prova amplificatore",
                "operation_id": "quick-withdrawal-test-0001",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("Prelevati 3 pezzi".encode(), response.data)
        with self.app.app_context():
            db = get_db()
            self.assertEqual(
                db.execute("SELECT quantity FROM components WHERE id = 1").fetchone()[0],
                7,
            )
            movement = db.execute("SELECT * FROM movements").fetchone()
            self.assertEqual(movement["movement_type"], "prelievo")
            self.assertEqual(movement["quantity"], 3)
            self.assertEqual(movement["delta"], -3)
            self.assertEqual(movement["balance_after"], 7)
            self.assertEqual(movement["recipient"], "Prof.ssa Maria Rossi")
            self.assertEqual(
                movement["note"],
                "Destinazione: Classe 4A Elettronica — Prova amplificatore",
            )
            self.assertEqual(movement["user_id"], operator_id)
            self.assertEqual(movement["operation_id"], "quick-withdrawal-test-0001")

    def test_quick_withdrawal_operation_id_prevents_double_stock_change(self):
        operator_id = self.create_operator()
        self.use_account(operator_id)
        payload = {
            "movement_type": "prelievo",
            "quantity": "4",
            "destination": "Laboratorio TPS",
            "operation_id": "same-mobile-operation-0001",
        }

        first = self.client.post("/q/TPS0001", data=payload)
        duplicate = self.client.post(
            "/q/TPS0001", data=payload, follow_redirects=True
        )
        self.assertEqual(first.status_code, 302)
        self.assertEqual(duplicate.status_code, 200)
        self.assertIn("Operazione gi\u00e0 registrata".encode(), duplicate.data)
        with self.app.app_context():
            db = get_db()
            self.assertEqual(
                db.execute("SELECT quantity FROM components WHERE id = 1").fetchone()[0],
                6,
            )
            self.assertEqual(
                db.execute(
                    "SELECT COUNT(*) FROM movements WHERE operation_id = ?",
                    (payload["operation_id"],),
                ).fetchone()[0],
                1,
            )

    def test_quick_withdrawal_rejects_insufficient_stock(self):
        operator_id = self.create_operator()
        self.use_account(operator_id)

        response = self.client.post(
            "/q/TPS0001",
            data={
                "movement_type": "prelievo",
                "quantity": "11",
                "operation_id": "insufficient-stock-test-0001",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("Giacenza insufficiente: disponibili 10 pezzi".encode(), response.data)
        with self.app.app_context():
            db = get_db()
            self.assertEqual(
                db.execute("SELECT quantity FROM components WHERE id = 1").fetchone()[0],
                10,
            )
            self.assertEqual(db.execute("SELECT COUNT(*) FROM movements").fetchone()[0], 0)

    def test_operator_cannot_record_loads_or_stock_adjustments(self):
        operator_id = self.create_operator()
        self.use_account(operator_id)

        page = self.client.get("/movimenti/nuovo")
        self.assertNotIn(b'<option value="carico">', page.data)
        self.assertNotIn(b'<option value="rettifica_carico">', page.data)
        self.assertNotIn(b'<option value="rettifica_scarico">', page.data)
        for movement_type in ("carico", "rettifica_carico", "rettifica_scarico"):
            with self.subTest(movement_type=movement_type):
                response = self.client.post(
                    "/movimenti/nuovo",
                    data={
                        "component_id": "1",
                        "movement_type": movement_type,
                        "quantity": "5",
                        "operation_id": f"forbidden-{movement_type}-0001",
                    },
                )
                self.assertEqual(response.status_code, 200)
                self.assertIn(
                    "I docenti possono registrare solo prelievi e restituzioni".encode(),
                    response.data,
                )

        with self.app.app_context():
            db = get_db()
            self.assertEqual(
                db.execute("SELECT quantity FROM components WHERE id = 1").fetchone()[0],
                10,
            )
            self.assertEqual(db.execute("SELECT COUNT(*) FROM movements").fetchone()[0], 0)

    def test_operator_can_record_withdrawals_and_returns_from_generic_form(self):
        operator_id = self.create_operator()
        self.use_account(operator_id)

        withdrawal = self.client.post(
            "/movimenti/nuovo",
            data={
                "component_id": "1",
                "movement_type": "prelievo",
                "quantity": "2",
                "recipient": "Classe 4A",
                "operation_id": "generic-operator-withdrawal-0001",
            },
        )
        returned = self.client.post(
            "/movimenti/nuovo",
            data={
                "component_id": "1",
                "movement_type": "restituzione",
                "quantity": "1",
                "recipient": "Classe 4A",
                "operation_id": "generic-operator-return-0001",
            },
        )

        self.assertEqual(withdrawal.status_code, 302)
        self.assertEqual(returned.status_code, 302)
        with self.app.app_context():
            db = get_db()
            self.assertEqual(
                db.execute("SELECT quantity FROM components WHERE id = 1").fetchone()[0],
                9,
            )
            movements = db.execute(
                "SELECT movement_type, delta, user_id FROM movements ORDER BY id"
            ).fetchall()
            self.assertEqual(
                [(row["movement_type"], row["delta"], row["user_id"]) for row in movements],
                [
                    ("prelievo", -2, operator_id),
                    ("restituzione", 1, operator_id),
                ],
            )

    def test_qr_url_configuration_validates_input_and_requires_admin(self):
        invalid = self.client.post(
            "/etichette-qr",
            data={"public_base_url": "https://utente:password@example.edu"},
            follow_redirects=True,
        )
        self.assertEqual(invalid.status_code, 200)
        self.assertIn("Inserire un indirizzo completo".encode(), invalid.data)
        with self.app.app_context():
            self.assertIsNone(
                get_db().execute(
                    "SELECT value FROM app_settings WHERE key = 'public_base_url'"
                ).fetchone()
            )

        valid = self.client.post(
            "/etichette-qr",
            data={"public_base_url": "https://magazzino.itis.example/"},
            follow_redirects=True,
        )
        self.assertEqual(valid.status_code, 200)
        self.assertIn("Indirizzo dei QR salvato".encode(), valid.data)
        self.assertIn(
            b"https://magazzino.itis.example/q/TPS0001",
            valid.data,
        )
        with self.app.app_context():
            saved_url = get_db().execute(
                "SELECT value FROM app_settings WHERE key = 'public_base_url'"
            ).fetchone()[0]
            self.assertEqual(saved_url, "https://magazzino.itis.example")

        operator_id = self.create_operator()
        self.use_account(operator_id)
        self.assertEqual(self.client.get("/etichette-qr").status_code, 403)
        forbidden = self.client.post(
            "/etichette-qr",
            data={"public_base_url": "https://changed.example"},
        )
        self.assertEqual(forbidden.status_code, 403)
        with self.app.app_context():
            saved_url = get_db().execute(
                "SELECT value FROM app_settings WHERE key = 'public_base_url'"
            ).fetchone()[0]
            self.assertEqual(saved_url, "https://magazzino.itis.example")

    def test_qr_svg_and_label_pdfs_are_valid_downloads(self):
        self.client.post(
            "/etichette-qr",
            data={"public_base_url": "https://magazzino.itis.example"},
        )

        svg = self.client.get("/qr/TPS0001.svg")
        self.assertEqual(svg.status_code, 200)
        self.assertEqual(svg.mimetype, "image/svg+xml")
        root = ET.fromstring(svg.data)
        self.assertEqual(root.tag.rsplit("}", 1)[-1], "svg")
        self.assertTrue(root.findall(".//{http://www.w3.org/2000/svg}path"))

        labels = self.client.post(
            "/etichette-qr.pdf", data={"selected_ids": ["1"]}
        )
        self.assertEqual(labels.status_code, 200)
        self.assertEqual(labels.mimetype, "application/pdf")
        self.assertTrue(labels.data.startswith(b"%PDF"))
        self.assertGreater(len(labels.data), 1000)
        self.assertIn("attachment", labels.headers["Content-Disposition"])

        single_label = self.client.get("/componenti/1/etichetta.pdf")
        self.assertEqual(single_label.status_code, 200)
        self.assertEqual(single_label.mimetype, "application/pdf")
        self.assertTrue(single_label.data.startswith(b"%PDF"))

    def test_reorder_report_uses_target_minus_current_stock(self):
        self.client.post(
            "/movimenti/nuovo",
            data={
                "component_id": "1",
                "movement_type": "prelievo",
                "quantity": "7",
                "recipient": "Lab TPS",
            },
        )
        response = self.client.get("/riordino")
        self.assertIn(b"TPS0001", response.data)
        self.assertIn(b'name="order_quantity_1"', response.data)
        self.assertIn(b'value="7"', response.data)

    def test_report_warns_before_reaching_threshold(self):
        self.client.post(
            "/movimenti/nuovo",
            data={
                "component_id": "1",
                "movement_type": "prelievo",
                "quantity": "6",
                "recipient": "Lab TPS",
            },
        )
        response = self.client.get("/riordino")
        self.assertIn("articoli vicini alla soglia".encode(), response.data)
        self.assertIn(b"TPS0001", response.data)

    def test_zero_stock_without_target_appears_in_reorder_report(self):
        with self.app.app_context():
            db = get_db()
            db.execute(
                """
                INSERT INTO components
                    (code, description, quantity, reorder_threshold, target_quantity)
                VALUES ('TPS0002', 'Componente esaurito', 0, 0, 0)
                """
            )
            db.execute(
                """
                INSERT INTO components
                    (code, description, quantity, reorder_threshold, target_quantity)
                VALUES ('TPS0003', 'Obiettivo ancora da definire', 2, 3, 0)
                """
            )
            db.commit()
        response = self.client.get("/riordino")
        self.assertIn(b"TPS0002", response.data)
        self.assertIn(b"TPS0003", response.data)
        self.assertIn(b"Da definire", response.data)

        excel = self.client.get("/riordino.xlsx")
        self.assertEqual(excel.status_code, 200)
        workbook = load_workbook(BytesIO(excel.data), data_only=True)
        rows = list(workbook["Distinta da completare"].iter_rows(values_only=True))
        zero_item = next(row for row in rows if row[0] == "TPS0002")
        self.assertEqual(zero_item[5], "DA DEFINIRE")
        self.assertEqual(zero_item[6], "DA DEFINIRE")

    def test_bulk_order_quantity_applies_equally_to_selected_components(self):
        with self.app.app_context():
            db = get_db()
            db.execute(
                """INSERT INTO components
                   (code, description, quantity, reorder_threshold, target_quantity)
                   VALUES ('TPS0002', 'Esaurito', 0, 0, 0)"""
            )
            db.execute(
                """INSERT INTO components
                   (code, description, quantity, reorder_threshold, target_quantity)
                   VALUES ('TPS0003', 'Sotto soglia', 2, 3, 5)"""
            )
            db.commit()

        response = self.client.post(
            "/riordino/quantita",
            data={
                "action": "bulk",
                "bulk_order_quantity": "4",
                "selected_ids": ["2", "3"],
            },
            follow_redirects=True,
        )
        self.assertIn("Aggiunti al modulo 2 componenti".encode(), response.data)
        with self.app.app_context():
            db = get_db()
            rows = db.execute(
                "SELECT code, quantity, target_quantity FROM components ORDER BY code"
            ).fetchall()
            self.assertEqual(rows[0]["target_quantity"], 10)
            self.assertEqual(rows[1]["target_quantity"], 0)
            self.assertEqual(rows[2]["target_quantity"], 5)
            order_lines = db.execute(
                """
                SELECT c.code, r.desired_quantity
                FROM reorder_lines r
                JOIN components c ON c.id = r.component_id
                ORDER BY c.code
                """
            ).fetchall()
            self.assertEqual(
                [(row["code"], row["desired_quantity"]) for row in order_lines],
                [("TPS0002", 4), ("TPS0003", 4)],
            )

        pending = self.client.get("/riordino")
        self.assertNotIn(b"TPS0002", pending.data)
        self.assertNotIn(b"TPS0003", pending.data)

        order_form = self.client.get("/modulo-riordino")
        self.assertIn(b"TPS0002", order_form.data)
        self.assertIn(b"TPS0003", order_form.data)
        self.assertIn(b"Pezzi totali", order_form.data)

        excel = self.client.get("/modulo-riordino.xlsx")
        workbook = load_workbook(BytesIO(excel.data), data_only=True)
        exported = {
            row[0]: row[4]
            for row in workbook["Modulo riordino"].iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(exported["TPS0002"], 4)
        self.assertEqual(exported["TPS0003"], 4)

        pending_excel = self.client.get("/riordino.xlsx")
        pending_workbook = load_workbook(BytesIO(pending_excel.data), data_only=True)
        pending_codes = {
            row[0]
            for row in pending_workbook["Distinta da completare"].iter_rows(
                min_row=2, values_only=True
            )
        }
        self.assertNotIn("TPS0002", pending_codes)
        self.assertNotIn("TPS0003", pending_codes)

    def test_individual_order_quantity_can_differ_from_threshold_and_return_to_pending(self):
        with self.app.app_context():
            db = get_db()
            db.execute(
                """INSERT INTO components
                   (code, description, quantity, reorder_threshold, target_quantity)
                   VALUES ('TPS0002', 'Esaurito', 0, 6, 0)"""
            )
            db.commit()

        added = self.client.post(
            "/riordino/quantita",
            data={"action": "single:2", "order_quantity_2": "5"},
            follow_redirects=True,
        )
        self.assertIn("TPS0002 aggiunto al modulo con quantit".encode(), added.data)
        with self.app.app_context():
            row = get_db().execute(
                """
                SELECT c.target_quantity, r.desired_quantity
                FROM components c
                JOIN reorder_lines r ON r.component_id = c.id
                WHERE c.code = 'TPS0002'
                """
            ).fetchone()
            self.assertEqual(row["target_quantity"], 0)
            self.assertEqual(row["desired_quantity"], 5)

        updated = self.client.post(
            "/modulo-riordino/2/quantita",
            data={"desired_quantity": "8"},
            follow_redirects=True,
        )
        self.assertIn("modulo aggiornata".encode(), updated.data)
        with self.app.app_context():
            self.assertEqual(
                get_db().execute(
                    "SELECT desired_quantity FROM reorder_lines WHERE component_id = 2"
                ).fetchone()[0],
                8,
            )

        removed = self.client.post(
            "/modulo-riordino/2/rimuovi", follow_redirects=True
        )
        self.assertIn("riportato nella distinta".encode(), removed.data)
        with self.app.app_context():
            db = get_db()
            self.assertIsNone(
                db.execute(
                    "SELECT 1 FROM reorder_lines WHERE component_id = 2"
                ).fetchone()
            )
            self.assertEqual(
                db.execute(
                    "SELECT target_quantity FROM components WHERE id = 2"
                ).fetchone()[0],
                0,
            )
        self.assertIn(b"TPS0002", self.client.get("/riordino").data)

    def test_operator_can_view_but_cannot_modify_order_form(self):
        with self.app.app_context():
            db = get_db()
            cursor = db.execute(
                "INSERT INTO users (username, full_name, password_hash, role) VALUES (?, ?, ?, ?)",
                ("operatore", "Operatore Test", generate_password_hash("password-test"), "operatore"),
            )
            operator_id = cursor.lastrowid
            component = db.execute(
                """INSERT INTO components
                   (code, description, quantity, reorder_threshold, target_quantity)
                   VALUES ('TPS0002', 'Esaurito', 0, 0, 0)"""
            )
            component_id = component.lastrowid
            db.execute(
                "INSERT INTO reorder_lines (component_id, desired_quantity) VALUES (?, 6)",
                (component_id,),
            )
            db.commit()
        with self.client.session_transaction() as session_data:
            session_data["user_id"] = operator_id
            session_data["auth_version"] = 1

        page = self.client.get("/modulo-riordino")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"TPS0002", page.data)
        self.assertNotIn(b'name="desired_quantity"', page.data)
        self.assertNotIn("Rimetti in distinta".encode(), page.data)
        self.assertEqual(self.client.get("/modulo-riordino.xlsx").status_code, 200)
        self.assertEqual(self.client.get("/modulo-riordino.pdf").status_code, 200)

        add_response = self.client.post(
            "/riordino/quantita",
            data={"action": "single:1", "order_quantity_1": "5"},
        )
        update_response = self.client.post(
            f"/modulo-riordino/{component_id}/quantita",
            data={"desired_quantity": "9"},
        )
        remove_response = self.client.post(
            f"/modulo-riordino/{component_id}/rimuovi"
        )
        self.assertEqual(add_response.status_code, 403)
        self.assertEqual(update_response.status_code, 403)
        self.assertEqual(remove_response.status_code, 403)
        with self.app.app_context():
            self.assertEqual(
                get_db().execute(
                    "SELECT desired_quantity FROM reorder_lines WHERE component_id = ?",
                    (component_id,),
                ).fetchone()[0],
                6,
            )

    def test_legacy_target_migration_is_idempotent_and_preserves_existing_lines(self):
        with self.app.app_context():
            db = get_db()
            db.execute(
                "DELETE FROM app_settings WHERE key = 'reorder_lines_migrated_from_targets_v1'"
            )
            migrated_component = db.execute(
                """INSERT INTO components
                   (code, description, quantity, reorder_threshold, target_quantity)
                   VALUES ('TPS0002', 'Legacy da migrare', 0, 0, 12)"""
            ).lastrowid
            existing_component = db.execute(
                """INSERT INTO components
                   (code, description, quantity, reorder_threshold, target_quantity)
                   VALUES ('TPS0003', 'Riga esistente', 0, 0, 9)"""
            ).lastrowid
            positive_component = db.execute(
                """INSERT INTO components
                   (code, description, quantity, reorder_threshold, target_quantity)
                   VALUES ('TPS0004', 'Legacy sotto soglia', 2, 3, 10)"""
            ).lastrowid
            natural_depletion = db.execute(
                """INSERT INTO components
                   (code, description, quantity, reorder_threshold, target_quantity,
                    created_at, updated_at)
                   VALUES ('TPS0005', 'Esaurimento naturale', 0, 0, 10,
                           '2026-01-01 08:00:00', '2026-01-02 08:00:00')"""
            ).lastrowid
            confirmed_after_movement = db.execute(
                """INSERT INTO components
                   (code, description, quantity, reorder_threshold, target_quantity,
                    created_at, updated_at)
                   VALUES ('TPS0006', 'Conferma dopo prelievo', 2, 3, 6,
                           '2026-01-01 08:00:00', '2026-01-03 08:00:00')"""
            ).lastrowid
            db.executemany(
                """INSERT INTO movements
                   (component_id, movement_type, quantity, delta, balance_after,
                    recipient, user_id, created_at)
                   VALUES (?, 'prelievo', ?, ?, ?, 'Laboratorio', 1, ?)""",
                [
                    (natural_depletion, 10, -10, 0, "2026-01-02 08:00:00"),
                    (confirmed_after_movement, 8, -8, 2, "2026-01-02 08:00:00"),
                ],
            )
            db.execute(
                "INSERT INTO reorder_lines (component_id, desired_quantity) VALUES (?, 4)",
                (existing_component,),
            )

            migrate_reorder_lines(db)
            rows = db.execute(
                """
                SELECT c.code, r.desired_quantity
                FROM reorder_lines r
                JOIN components c ON c.id = r.component_id
                ORDER BY c.code
                """
            ).fetchall()
            self.assertEqual(
                [(row["code"], row["desired_quantity"]) for row in rows],
                [
                    ("TPS0002", 12),
                    ("TPS0003", 4),
                    ("TPS0004", 8),
                    ("TPS0006", 4),
                ],
            )
            self.assertIsNone(
                db.execute(
                    "SELECT 1 FROM reorder_lines WHERE component_id = ?",
                    (natural_depletion,),
                ).fetchone()
            )
            self.assertEqual(
                db.execute(
                    "SELECT desired_quantity FROM reorder_lines WHERE component_id = ?",
                    (positive_component,),
                ).fetchone()[0],
                8,
            )

            db.execute(
                "UPDATE components SET target_quantity = 20 WHERE id = ?",
                (migrated_component,),
            )
            migrate_reorder_lines(db)
            rows_after_second_run = db.execute(
                "SELECT component_id, desired_quantity FROM reorder_lines ORDER BY component_id"
            ).fetchall()
            self.assertEqual(len(rows_after_second_run), 4)
            self.assertEqual(
                db.execute(
                    "SELECT desired_quantity FROM reorder_lines WHERE component_id = ?",
                    (migrated_component,),
                ).fetchone()[0],
                12,
            )
            self.assertIsNotNone(
                db.execute(
                    """SELECT 1 FROM app_settings
                       WHERE key = 'reorder_lines_migrated_from_targets_v1'"""
                ).fetchone()
            )

    def test_legacy_movement_schema_adds_operation_id_without_losing_history(self):
        legacy_path = str(Path(self.temp.name) / "legacy.db")
        legacy_db = sqlite3.connect(legacy_path)
        legacy_db.execute(
            """
            CREATE TABLE movements (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                component_id INTEGER NOT NULL,
                movement_type TEXT NOT NULL,
                quantity INTEGER NOT NULL,
                delta INTEGER NOT NULL,
                balance_after INTEGER NOT NULL,
                recipient TEXT,
                note TEXT,
                user_id INTEGER NOT NULL,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        legacy_db.execute(
            """INSERT INTO movements
               (component_id, movement_type, quantity, delta, balance_after,
                recipient, note, user_id)
               VALUES (7, 'prelievo', 2, -2, 8, 'Prof. Legacy', 'Storico', 3)"""
        )
        legacy_db.commit()
        legacy_db.close()

        legacy_config = {
            "TESTING": True,
            "DATABASE": legacy_path,
            "EXCEL_SOURCE": str(Path(self.temp.name) / "missing-legacy.xlsx"),
            "CSRF_ENABLED": False,
            "SECRET_KEY": "legacy-test",
        }
        legacy_app = create_app(legacy_config)
        with legacy_app.app_context():
            db = get_db()
            columns = {
                row["name"] for row in db.execute("PRAGMA table_info(movements)")
            }
            self.assertIn("operation_id", columns)
            movement = db.execute("SELECT * FROM movements WHERE id = 1").fetchone()
            self.assertEqual(movement["recipient"], "Prof. Legacy")
            self.assertEqual(movement["balance_after"], 8)
            self.assertIsNone(movement["operation_id"])
            indexes = {
                row["name"] for row in db.execute("PRAGMA index_list(movements)")
            }
            self.assertIn("idx_movements_operation_id", indexes)

        second_start = create_app(legacy_config)
        with second_start.app_context():
            db = get_db()
            columns = [
                row["name"] for row in db.execute("PRAGMA table_info(movements)")
            ]
            self.assertEqual(columns.count("operation_id"), 1)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM movements").fetchone()[0], 1)

    def test_movement_history_is_immutable_in_database(self):
        self.client.post(
            "/movimenti/nuovo",
            data={
                "component_id": "1",
                "movement_type": "carico",
                "quantity": "1",
                "recipient": "",
            },
        )
        with self.app.app_context():
            with self.assertRaises(sqlite3.IntegrityError):
                get_db().execute("DELETE FROM movements")

    def test_exports_are_downloadable(self):
        self.client.post(
            "/movimenti/nuovo",
            data={
                "component_id": "1",
                "movement_type": "prelievo",
                "quantity": "7",
                "recipient": "Lab TPS",
            },
        )
        excel = self.client.get("/riordino.xlsx")
        pdf = self.client.get("/riordino.pdf")
        self.assertEqual(excel.status_code, 200)
        self.assertTrue(excel.data.startswith(b"PK"))
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.data.startswith(b"%PDF"))


if __name__ == "__main__":
    unittest.main()
